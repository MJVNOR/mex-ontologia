"""Inspección reanudable de archivos y diccionarios; no conserva filas de datos."""
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
import csv
import hashlib
import io
import json
from pathlib import Path
import sys
import re
from urllib.request import Request, urlopen
from zipfile import ZipFile

import openpyxl
import xlrd
from inventariar_portal import ROOT, api, normalizar

CACHE = ROOT / 'data/portal/inspecciones'
MAX = 20 * 1024 * 1024


class ArchivoRemoto(io.RawIOBase):
    """Acceso seekable por rangos para ZIP/XLSX; caché en memoria con límite."""
    def __init__(self, url):
        self.url, self.pos, self.cache, self.bytes_leidos = url, 0, {}, 0
        with urlopen(Request(url, headers={'Range': 'bytes=0-1', 'User-Agent': 'MEX-Ontologi/0.2'}), timeout=40) as r:
            match = re.fullmatch(r'bytes 0-1/(\d+)', r.headers.get('Content-Range', ''))
            if r.status != 206 or not match:
                raise ValueError('Servidor no permite acceso parcial verificable')
            self.size = int(match[1])
            self.etag = r.headers.get('ETag')
            r.read(2)

    def seekable(self):
        return True

    def readable(self):
        return True

    def tell(self):
        return self.pos

    def seek(self, offset, whence=0):
        self.pos = offset + (self.pos if whence == 1 else self.size if whence == 2 else 0)
        if self.pos < 0:
            raise ValueError('Posición negativa')
        return self.pos

    def read(self, size=-1):
        end = self.size if size < 0 else min(self.size, self.pos + size)
        parts, chunk = [], 262144
        while self.pos < end:
            start = (self.pos // chunk) * chunk
            if start not in self.cache:
                stop = min(start + chunk, self.size) - 1
                if self.bytes_leidos + stop - start + 1 > 64 * 1024 * 1024:
                    raise ValueError('Acceso remoto excede 64 MiB')
                with urlopen(Request(self.url, headers={'Range': f'bytes={start}-{stop}', 'User-Agent': 'MEX-Ontologi/0.2'}), timeout=40) as r:
                    if r.status != 206 or r.headers.get('Content-Range') != f'bytes {start}-{stop}/{self.size}' or (self.etag and r.headers.get('ETag') != self.etag):
                        raise ValueError('Rango inválido o archivo cambió')
                    block = r.read(stop - start + 2)
                    if len(block) != stop - start + 1:
                        raise ValueError('Rango incompleto')
                self.cache[start] = block
                self.bytes_leidos += len(block)
            length = min(end - self.pos, start + len(self.cache[start]) - self.pos)
            parts.append(self.cache[start][self.pos - start:self.pos - start + length])
            self.pos += length
        return b''.join(parts)


def tablas_xlsx_remoto(url):
    remote = ArchivoRemoto(url)
    book = openpyxl.load_workbook(remote, read_only=True, data_only=True)
    try:
        sheets = [(s.title, list_rows(s.iter_rows(values_only=True), 30)) for s in book.worksheets]
    finally:
        book.close()
    return sheets, {'metodo': 'xlsx_rangos_http', 'bytes_leidos': remote.bytes_leidos,
                     'tamano_archivo': remote.size, 'etag': remote.etag, 'completo': False}


def descargar(url, parcial=False):
    req = Request(url, headers={'User-Agent': 'MEX-Ontologi/0.2 (public data structure inspection)',
                               **({'Range': 'bytes=0-131071'} if parcial else {})})
    with urlopen(req, timeout=40) as response:
        limit = 131072 if parcial else MAX
        data = response.read(limit + 1)
        completo = len(data) <= limit and response.status != 206
        if not parcial and len(data) > limit:
            raise ValueError('Archivo excede límite de 20 MiB')
        return data[:limit], {'url_final': response.url, 'http': response.status,
                              'bytes_leidos': min(len(data), limit), 'completo': completo,
                              'sha256_fragmento_o_archivo': hashlib.sha256(data[:limit]).hexdigest()}


def tablas(data, formato, completo, diccionario=False):
    """Hasta 30 filas para detectar encabezado; diccionarios tienen otro límite de lectura."""
    if data.startswith(b'PK'):
        with ZipFile(io.BytesIO(data)) as z:
            if sum(f.file_size for f in z.infolist()) > 150 * 1024 * 1024:
                raise ValueError('Excel expandido excede límite')
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            return [(sheet.title, list_rows(sheet.iter_rows(values_only=True), 10001 if diccionario else 30)) for sheet in book.worksheets]
        finally:
            book.close()
    if data.startswith(bytes.fromhex('D0CF11E0')):
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            return [(s.name, [s.row_values(i) for i in range(min(s.nrows, 10001 if diccionario else 30))]) for s in book.sheets()]
        finally:
            book.release_resources()
    for encoding in ('utf-8-sig', 'cp1252', 'latin1'):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            pass
    if text.lstrip().lower().startswith(('<!doctype html', '<html')):
        raise ValueError('URL devolvió HTML en lugar de archivo')
    if formato == 'JSON':
        value = json.loads(text)
        if isinstance(value, dict) and value.get('type') == 'FeatureCollection':
            return [('GeoJSON:properties', [list(value['features'][0].get('properties', {}))])]
        rows = value if isinstance(value, list) else value.get('records', [])
        if rows and isinstance(rows[0], dict):
            return [('JSON', [list(rows[0])])]
        raise ValueError('JSON sin tabla reconocida')
    if not completo:
        text = text.rsplit('\n', 1)[0]
    try:
        dialect = csv.Sniffer().sniff(text[:10000], delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel
    return [('CSV', list_rows(csv.reader(io.StringIO(text), dialect), 10001 if diccionario else 30))]


def list_rows(iterator, limit=30):
    rows = []
    for row in iterator:
        rows.append([str(v) if v is not None else '' for v in row])
        if len(rows) >= limit:
            break
    return rows


def encabezado(rows):
    candidates = []
    for i, row in enumerate(rows[:30]):
        cells = [str(c).strip() for c in row]
        nonempty = [c for c in cells if c]
        if not nonempty:
            continue
        dictionary = any('nombre' in normalizar(c) and ('variable' in normalizar(c) or 'campo' in normalizar(c)) for c in cells)
        score = len(nonempty) + (100 if dictionary else 0)
        # Headers with text are preferred over dense numeric data rows.
        score += sum(any(c.isalpha() for c in value) for value in nonempty)
        candidates.append((score, -i, cells))
    if not candidates:
        raise ValueError('Hoja sin encabezado observable')
    _, neg, cells = max(candidates)
    fields = [{'id': c, 'posicion_original': i + 1} for i, c in enumerate(cells) if c]
    return -neg, fields


def inspeccionar(r):
    path = CACHE / (r['id'] + '.json')
    if path.exists():
        cached = json.loads(path.read_text(encoding='utf-8'))
        reuse = cached.get('url_recurso_al_inspeccionar', r.get('url')) == r.get('url')
        if '--refresh' not in sys.argv and reuse and (cached['ok'] or '--retry-errors' not in sys.argv):
            return cached
    result = {'id': r['id'], 'consultado_en_utc': datetime.now(timezone.utc).isoformat(),
              'fuente': r.get('url'), 'url_recurso_al_inspeccionar': r.get('url'), 'ok': False, 'tablas': []}
    dictionary = r['clasificacion'] == 'diccionario'
    try:
        if dictionary and r.get('datastore', {}).get('ok'):
            records, offset = [], 0
            while True:
                response = api('datastore_search', resource_id=r['id'], limit=1000, offset=offset, sort='_id asc')
                if not response['ok']:
                    raise ValueError(str(response.get('error')))
                body = response['result']
                batch = body.get('records', [])
                records.extend({k: v for k, v in row.items() if not k.startswith('_')} for row in batch)
                offset += len(batch)
                if offset >= body.get('total', 0):
                    break
                if not batch or offset >= 10000:
                    raise ValueError('Diccionario incompleto por límite o paginación')
            result.update(ok=True, metodo='api_diccionario', fuente=response['url'], diccionario_completo=True,
                          definiciones=records)
        elif r.get('datastore', {}).get('ok') and not dictionary:
            result.update(ok=True, metodo='api_esquema_previo', fuente=r['datastore']['url'])
        else:
            if not r.get('url'):
                metadata = api('resource_show', id=r['id'])
                datastore = api('datastore_search', resource_id=r['id'], limit=0)
                result['comprobaciones_sin_url'] = {
                    'resource_show_url': metadata.get('url'),
                    'url_publicada_actual': metadata.get('result', {}).get('url'),
                    'datastore': {k: v for k, v in datastore.items() if k != 'result'}}
                raise ValueError('Sin URL')
            try:
                data, evidence = descargar(r['url'], parcial=r['formato'].upper() == 'CSV' and not dictionary)
                result.update(evidence)
                result['metodo'] = 'archivo_acotado_en_memoria'
                sheets = tablas(data, r['formato'].upper(), evidence['completo'], dictionary)
            except ValueError as error:
                if dictionary or 'límite' not in str(error):
                    raise
                if r['formato'].upper() == 'XLSX':
                    sheets, evidence = tablas_xlsx_remoto(r['url'])
                    result.update(evidence)
                else:
                    # BIFF/XLS requiere el contenedor completo para localizar hojas.
                    with urlopen(Request(r['url'], headers={'User-Agent': 'MEX-Ontologi/0.2'}), timeout=40) as response:
                        data = response.read(64 * 1024 * 1024 + 1)
                    if len(data) > 64 * 1024 * 1024:
                        raise ValueError('XLS excede límite de 64 MiB')
                    sheets = tablas(data, 'XLS', True)
                    result.update(metodo='xls_completo_acotado_en_memoria', bytes_leidos=len(data), completo=True,
                                  sha256_fragmento_o_archivo=hashlib.sha256(data).hexdigest())
            definitions = []
            for name, rows in sheets:
                if not any(any(str(v).strip() for v in row) for row in rows):
                    continue
                start, fields = encabezado(rows)
                result['tablas'].append({'hoja': name, 'fila_encabezado': start + 1,
                                         'campos': fields, 'metodo_encabezado': 'heuristica_requiere_revision'})
                if dictionary:
                    if len(rows) >= 10001:
                        raise ValueError('Diccionario excede límite de filas')
                    for row in rows[start + 1:]:
                        if any(str(v).strip() for v in row):
                            definitions.append({'hoja': name, **{f['id']: row[f['posicion_original'] - 1]
                                                if len(row) >= f['posicion_original'] else '' for f in fields}})
            if not result['tablas']:
                raise ValueError('Sin tablas reconocidas')
            result['ok'] = True
            if dictionary:
                result.update(definiciones=definitions, diccionario_completo=evidence['completo'])
    except Exception as error:
        result.update(ok=False, error=f'{type(error).__name__}: {error}')
    path.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    return result


def main():
    CACHE.mkdir(parents=True, exist_ok=True)
    inv = json.loads((ROOT / 'data/portal/inventario.json').read_text(encoding='utf-8'))
    results = []
    with ThreadPoolExecutor(3) as pool:
        jobs = [pool.submit(inspeccionar, r) for r in inv['recursos']]
        for job in as_completed(jobs):
            results.append(job.result())
            if len(results) % 20 == 0:
                print(f'Inspeccionados {len(results)}/{len(jobs)}; correctos {sum(r["ok"] for r in results)}', flush=True)
    report = {'generado_en_utc': datetime.now(timezone.utc).isoformat(), 'recursos': sorted(results, key=lambda r: r['id'])}
    (ROOT / 'data/portal/contenido-inspeccion.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print('Finalizado', len(results), 'correctos', sum(r['ok'] for r in results), flush=True)


if __name__ == '__main__':
    main()
