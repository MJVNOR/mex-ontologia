"""Verificacion independiente CLASIF (auditoria-clasif). Solo lee; escribe resultado.json aqui mismo."""
import hashlib, json, struct, zipfile
from collections import Counter
from pathlib import Path
ROOT = Path(__file__).resolve().parents[3]
EV = ROOT / "propuestas/investigacion-post-r6/diversos/evidencia"
OUT = Path(__file__).resolve().parent / "resultado.json"
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
res = {"shas":{},"checks":[]}
def check(nombre, ok, detalle=""):
    res["checks"].append({"nombre":nombre,"ok":bool(ok),"detalle":detalle})
    return bool(ok)
# 1. SHAs esperados (de entrega.json)
esperados = {
 "propuestas/investigacion-post-r6/diversos/evidencia/sonora-imc2020-local.xlsx":"14a741bb438adbf82aa535baaf87abe0fab94fd23c8df86b84fa98d1904b5224",
 "propuestas/investigacion-post-r6/diversos/evidencia/conapo-IMC_2020.zip":"5d93b314428f88008ba380bac56b587169cfefd3acaf4659857a07874822b8bb",
 "propuestas/investigacion-post-r6/diversos/evidencia/conapo-imc2020_shp.zip":"69d73b5bb0127df6c81701684eb28cd6bd666d9319a1ffb45da76446ae278426",
 "propuestas/investigacion-post-r6/diversos/evidencia/ckan-marginacion.json":"44594c502e004937eb2a910815f1225cb1adffaa0fd226321dedbec14663a777",
 "propuestas/investigacion-post-r6/diversos/evidencia/inegi-rnm-v450.html":"0eb78483fa130f553b607efe792ef92f217725dd65873cae8d21596e054e9466",
 "propuestas/investigacion-post-r6/diversos/evidencia/conapo-imc-diccionario-nota.md":"97ac6f973389087eac54346b1893dd7cacaa8fba443b21bc14e2d7980f374ccd",
 "propuestas/investigacion-post-r6/diversos/evidencia/conapo-clasif-comparativa-nota.md":"e0f3d75a58418b3a7118b94027d1e2425c398b8b75702a5e9d019ce84a0352c4",
}
for ruta, esp in esperados.items():
    h = sha(ROOT/ruta); res["shas"][ruta]=h
    check("sha:"+Path(ruta).name, h==esp, h[:16])
# 2. Inspecciones locales existen
for iid in ["67156807-2013-4117-9ca9-394aea647127","09ef2342-fd9f-4334-bcd5-9c969baca32f","483259a2-b688-4f4b-beae-9c7016881a63"]:
    p = ROOT/f"data/portal/inspecciones/{iid}.json"
    d = json.loads(p.read_text(encoding="utf-8"))
    res["shas"][f"inspeccion:{iid}"]=sha(p)
    check("inspeccion:"+iid, p.is_file())
# diccionario local fila CLASIF
dloc = json.loads((ROOT/"data/portal/inspecciones/67156807-2013-4117-9ca9-394aea647127.json").read_text(encoding="utf-8"))
fila = [x for x in dloc["definiciones"] if x.get("Nombre variable")=="CLASIF"][0]
res["fila_local_clasif"]=fila
check("local_descripcion_generica", fila["Descripción"].strip().lower()=="clasificación", fila["Descripción"])
check("local_unidad_7_sin_catalogo", fila["Unidad de medida"].count(";")==6 and fila["Catálogo de referencia"]=="", fila["Unidad de medida"])
check("local_nulo_guion", fila["Identificador de valures nulos"]=="-", fila["Identificador de valures nulos"])
# recurso local hoja/columna
drec = json.loads((ROOT/"data/portal/inspecciones/09ef2342-fd9f-4334-bcd5-9c969baca32f.json").read_text(encoding="utf-8"))
check("recurso_hoja_IMC2020", drec["tablas"][0]["hoja"]=="IMC_2020")
check("recurso_clasif_pos5", [c for c in drec["tablas"][0]["campos"] if c["id"]=="CLASIF"][0]["posicion_original"]==5)
check("recurso_28_columnas", len(drec["tablas"][0]["campos"])==28, str(len(drec["tablas"][0]["campos"])))
dcat = json.loads((ROOT/"data/portal/inspecciones/483259a2-b688-4f4b-beae-9c7016881a63.json").read_text(encoding="utf-8"))
hojas = [t["hoja"] for t in dcat["tablas"]]
check("catalogo_sin_clasif", not any("clasif" in h.lower() for h in hojas), str(hojas))
# CKAN: fuentes extras a nivel conjunto, sin fila
ck = json.loads((EV/"ckan-marginacion.json").read_text(encoding="utf-8"))
res["ckan_fuentes_extras"]=ck["result"]["fuentes_extras"]
check("ckan_fuente_conjunto_no_fila", ck["result"]["fuentes_extras"]=="https://www.gob.mx/conapo/documentos/indices-de-marginacion-2020-284372", ck["result"]["fuentes_extras"])
# 3. Igualdad 2129 por clave y 28 campos (solo normalizacion documentada: strip texto; float==float)
from openpyxl import load_workbook
wbL = load_workbook(EV/"sonora-imc2020-local.xlsx", read_only=True, data_only=True)
rowsL = list(wbL[wbL.sheetnames[0]].iter_rows(values_only=True))
headL = [str(v) for v in rowsL[0]]
dL = {str(r[0]).strip(): r for r in rowsL[1:]}
z = zipfile.ZipFile(EV/"conapo-IMC_2020.zip")
tmp = Path(__file__).resolve().parent / "_tmp_imuc.xlsx"
tmp.write_bytes(z.read("IMUC_2020.xlsx"))
wbN = load_workbook(tmp, read_only=True, data_only=True)
dN = {}
for name in ["IMUC_2020_AGS-MOR","IMUC_2020_NAY-ZAC"]:
    ws = wbN[name]
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i==0:
            headN=[str(v) for v in r]; continue
        if str(r[5]).strip()=="26": dN[str(r[0]).strip()]=r
check("cabecera_28_identica", headL==headN, f"{len(headL)} cols")
def norm(v):
    if v is None: return ""
    if isinstance(v,float): return repr(v)
    return str(v).strip()
faltan = [k for k in dL if k not in dN]; sobran=[k for k in dN if k not in dL]
check("n_2129_ambos", len(dL)==2129 and len(dN)==2129, f"local={len(dL)} nac-son={len(dN)}")
check("claves_identicas", not faltan and not sobran, f"faltan={len(faltan)} sobran={len(sobran)}")
dif=0; ej=[]
for k in dL:
    rl, rn = dL[k], dN[k]
    for j in range(28):
        a,b = norm(rl[j]), norm(rn[j])
        if a!=b:
            try:
                if float(a)==float(b): continue
            except: pass
            dif+=1; ej.append({"clave":k,"col":headL[j],"local":repr(rl[j])[:80],"nacional":repr(rn[j])[:80]}); break
check("28_campos_por_clave_sin_dif", dif==0, f"diferencias={dif} ej={ej[:3]}")
res["distribucion_clasif_local"]=dict(Counter(norm(r[4]) for r in dL.values()))
res["distribucion_clasif_nacson"]=dict(Counter(norm(r[4]) for r in dN.values()))
check("distribucion_clasif_identica", res["distribucion_clasif_local"]==res["distribucion_clasif_nacson"], str(res["distribucion_clasif_local"]))
wsD = wbN["Diccionario"]
dic = [(r[0],r[1]) for r in wsD.iter_rows(values_only=True)]
res["diccionario_nacional_clasif"]=[list(x) for x in dic if x[0]=="CLASIF"]
check("diccionario_nacional_generico", res["diccionario_nacional_clasif"]==[["CLASIF","Clasificación "]], str(res["diccionario_nacional_clasif"]))
# 4. DBF: campo CLASIF sin definicion semantica; subset Sonora
raw = zipfile.ZipFile(EV/"conapo-imc2020_shp.zip").read("colonias_imc2020.dbf")
nreg,hlen,rlen = struct.unpack("<IHH", raw[4:12])
off=32; fields=[]
while raw[off]!=0x0D:
    fields.append((raw[off:off+11].split(b"\x00")[0].decode("latin-1"), raw[off+16])); off+=32
pos={n:(1+sum(l for _,l in fields[:i]),l) for i,(n,l) in enumerate(fields)}
res["dbf"]={"registros":nreg,"campos":len(fields),"nombres":[n for n,_ in fields]}
pe,le=pos["CVE_ENT"]; pc,lc=pos["CLASIF"]; pv,lv=pos["CVE_COL"]
cN=Counter(); nson=0
for i in range(nreg):
    b=hlen+i*rlen
    if raw[b+pe:b+pe+le].decode("latin-1").strip()=="26":
        nson+=1; cN[raw[b+pc:b+pc+lc].decode("latin-1").strip()]+=1
res["dbf_sonora_n"]=nson; res["dbf_sonora_dist"]=dict(cN)
check("dbf_sonora_2129", nson==2129, str(nson))
check("dbf_clasif_dist_identica", dict(cN)==res["distribucion_clasif_local"], str(dict(cN)))
cT=Counter()
for i in range(nreg):
    b=hlen+i*rlen; cT[raw[b+pc:b+pc+lc].decode("latin-1").strip()]+=1
res["dbf_nacional_distintos"]=len(cT); res["dbf_nacional_valores"]=sorted(cT)
check("dbf_38_observados_no_dominio", len(cT)==38, f"distintos={len(cT)}")
check("dbf_xlsx_conteo_difiere_1", nreg==74226, f"dbf={nreg} vs xlsx=74227 (nota: difieren en 1)")
# 5. Ausencia de afirmacion tipo-asentamiento en fuentes CONAPO examinadas
txt_inegi = (EV/"inegi-rnm-v450.html").read_text(encoding="utf-8", errors="replace")
check("inegi_otro_instrumento_sin_vinculo", "v_tip_asen" in txt_inegi and "Alojamientos de asistencia social" in txt_inegi, "RNM F4 alojamientos, no IMC_2020")
nota = (EV/"conapo-im2020-nota-webfetch.md").read_text(encoding="utf-8")
check("conapo_nota_sin_sentido_clasif", "tipo de asentamiento" not in nota.lower(), "nota solo acredita base a nivel colonia")
try: wbL.close()
except: pass
try: wbN.close()
except: pass
tmp.unlink(missing_ok=True)
res["ok_global"]=all(c["ok"] for c in res["checks"])
OUT.write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps({"ok_global":res["ok_global"],"fallos":[c for c in res["checks"] if not c["ok"]]}, ensure_ascii=False))
