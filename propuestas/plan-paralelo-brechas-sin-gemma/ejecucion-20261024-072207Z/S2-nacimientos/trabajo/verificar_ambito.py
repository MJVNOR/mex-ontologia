import os
import openpyxl
BASE = "propuestas/plan-paralelo-brechas-sin-gemma/ejecucion-20261024-072207Z/S2-nacimientos/trabajo/evidencia/cat_sinac_2020_2023"

def dump(f, want, maxrows=45):
    wb = openpyxl.load_workbook(os.path.join(BASE, f), read_only=True, data_only=True)
    ws = wb.active
    hit = []
    n = 0
    for r in ws.iter_rows(values_only=True):
        n += 1
        if r[0] is None and r[1] is None:
            continue
        s = " ".join(str(c) for c in r if c is not None)
        if any(w in s for w in want):
            hit.append(r)
        if len(hit) >= maxrows:
            break
    print("=====", f, "| filas:", ws.max_row)
    for h in hit:
        print("   ", h)
    if not hit:
        print("    SIN COINCIDENCIA")

dump("ENTIDADES.xlsx", ["26", "SONORA"])
dump("MUNICIPIOS_202201 (2).xlsx", ["SAN LUIS", "R\u00cdO COLORADO"])
dump("ESTABLECIMIENTO_SALUD_202204.xlsx", ["SRSMP000314"])
