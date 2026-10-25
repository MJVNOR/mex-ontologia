import glob, os
import openpyxl
BASE = "propuestas/plan-paralelo-brechas-sin-gemma/ejecucion-20261024-072207Z/S2-nacimientos/trabajo/evidencia/cat_sinac_2020_2023"
WANT = ("SI_NO", "TRIMESTRE", "AFILIACION", "OCUPACION", "PRODUCTO",
        "RESOLUCION", "UTILIZO", "PERSONA", "TIPO_MEDICO",
        "NACIMIENTO_CERT", "ENTIDADES")
for f in sorted(glob.glob(os.path.join(BASE, "*.xlsx"))):
    name = os.path.basename(f)
    if not name.startswith(WANT):
        continue
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = [r for r in ws.iter_rows(values_only=True)][:14]
        print("=====", name, "| hoja:", ws.title, "| filas_tot:",
              ws.max_row, "x", ws.max_column)
        for r in rows:
            print("   ", r)
