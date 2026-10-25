"""Materializa entregas revisadas desde copias fijadas por SHA-256, sin red.

Las propuestas originales no son dependencias de generación. Ninguna comprobación
de valores promociona una correspondencia semántica ni acredita vigencia histórica.
"""
from collections import Counter
import csv
import hashlib
import io
import json
import re
from urllib.parse import quote

from openpyxl import load_workbook
from rdflib import Graph, Literal, Namespace, OWL, RDF, RDFS, SKOS, URIRef, DCTERMS as DCT

from construir_catalogo import ROOT, PROV


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def composicion_exacta(clave, entidad):
    return (isinstance(clave, str) and isinstance(entidad, str)
            and re.fullmatch(r'[0-9]{5}', clave) is not None
            and re.fullmatch(r'[0-9]{2}', entidad) is not None and clave[:2] == entidad)


def leer_tabla(path, hoja=None):
    """Conserva textos de Excel; no convierte cadenas numéricas ni usa padding."""
    if path.suffix == '.xlsx':
        wb = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=False)
        try:
            ws = wb[hoja] if hoja else wb.worksheets[0]
            ws.reset_dimensions()
            rows = []
            for cells in ws.iter_rows():
                if not any(c.value is not None for c in cells):
                    continue
                if any(c.data_type == 'f' for c in cells):
                    raise ValueError('Fórmula sin revisión en fuente: ' + str(path))
                rows.append([c.value for c in cells])
        finally:
            wb.close()
        return rows
    raw = path.read_bytes()
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = raw.decode('latin1')
    return [r for r in csv.reader(io.StringIO(text)) if any(v.strip() for v in r)]


def texto(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def construir(semantic, semantic_data, inventory, inspection, base):
    from consolidar_semantica import columnas
    config_path = ROOT / 'ontologia/revision-agentes.json'
    config = json.loads(config_path.read_text(encoding='utf-8'))
    manifest_path = ROOT / config['manifiesto']
    if sha(manifest_path) != config['sha256_manifiesto']:
        raise ValueError('Manifiesto de entregas cambió; requiere revisión')
    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
    source_lookup = {k.replace('\\', '/'): v for k, v in manifest['fuentes'].items()}
    sources = {'revision_agentes': config_path, 'manifiesto_agentes': manifest_path}
    for group in ['entregas', 'fuentes']:
        for original, meta in manifest[group].items():
            path = ROOT / (original if group == 'entregas' else meta['archivo'])
            if not path.resolve().is_relative_to((ROOT / 'data/integracion').resolve()):
                raise ValueError('Evidencia fuera de la carpeta de integración')
            if sha(path) != meta['sha256']:
                raise ValueError('Evidencia de entrega cambió: ' + str(path))
            sources['entrega_' + str(path.relative_to(ROOT))] = path

    def propuesta(agent):
        return json.loads((ROOT / 'data/integracion/entregas' / agent / 'propuestas.json').read_text(encoding='utf-8'))

    def archivo(original):
        return ROOT / source_lookup[original.replace('\\', '/')]['archivo']

    mex, g, model = Namespace(base + 'vocab/'), Graph(), Graph()
    for graph in [g, model]:
        for prefix, ns in [('mex', mex), ('prov', PROV), ('skos', SKOS), ('dct', DCT)]:
            graph.bind(prefix, ns)

    def node(kind, key):
        return URIRef(base + kind + '/' + quote(str(key), safe=''))

    def lit(s, key, value):
        if value is not None:
            g.add((s, mex[key], Literal(value)))

    classes = {'CatalogoPublicado': 'Catálogo publicado identificado',
               'EntradaCatalogoPublicado': 'Entrada de catálogo publicado',
               'VinculoCatalogoPublicado': 'Vínculo documental a catálogo publicado',
               'ReferenciaLocalidadesObservada': 'Referencia de localidades observada en IML 2020',
               'LocalidadObservada': 'Localidad observada en IML 2020',
               'ComprobacionAplicacionTerritorial': 'Comprobación de claves territoriales'}
    objects = {'referenciaDeclarada': 'Referencia declarada', 'catalogoRecuperado': 'Catálogo recuperado',
               'variableDestino': 'Variable documental de destino', 'recursoComprobado': 'Recurso comprobado',
               'columnaComprobada': 'Columna física comprobada', 'referenciaComprobada': 'Referencia usada en comprobación',
               'referenciaLocalidadesObservada': 'Referencia de localidades observada en este recurso'}
    literals = {'codigoOriginal': 'Código original', 'nombreOriginal': 'Nombre original',
                'archivoSha256': 'SHA-256 del archivo publicado', 'hojaOriginal': 'Hoja original',
                'filaOriginal': 'Fila original', 'encabezadoOriginal': 'Encabezado original',
                'filaOriginalJson': 'Contenido literal de fila publicada (JSON)',
                'estadoVigencia': 'Estado de vigencia documental', 'estadoReferenciaObservada': 'Estado de referencia observada',
                'componenteLocalidad': 'Componente localidad', 'nombreMunicipioOriginal': 'Nombre original del municipio',
                'estadoComprobacion': 'Resultado de comprobación', 'coberturaComprobacion': 'Cobertura de comprobación',
                'campoComprobado': 'Nombre físico de campo comprobado', 'periodoComprobado': 'Periodo observado',
                'filasComprobadas': 'Filas comprobadas', 'claveObservada': 'Clave distinta observada',
                'claveEspecialNoMunicipal': 'Clave observada de categoría no municipal',
                'evidenciaComprobacionJson': 'Detalle reproducible de comprobación (JSON)',
                'sha256CatalogoComprobado': 'SHA-256 de catálogo usado en comprobación',
                'ubicacionComprobadaJson': 'Ubicación física comprobada (JSON)'}
    for key, label in classes.items():
        model.add((mex[key], RDF.type, OWL.Class)).add((mex[key], RDFS.label, Literal(label, lang='es')))
    for fields, cls in [(objects, OWL.ObjectProperty), (literals, OWL.DatatypeProperty)]:
        for key, label in fields.items():
            model.add((mex[key], RDF.type, cls)).add((mex[key], RDFS.label, Literal(label, lang='es')))
    variables = {v['definicion_iri']: v for v in semantic['variables']}
    resources = {r['id']: r for r in inventory['recursos']}
    for proposed in propuesta('agente-3-significados')['variables']:
        current = variables[proposed['definicion_iri']]
        if (current['huella_fila_diccionario'] != proposed['huella_fila_diccionario']
                or current['descripcion'] != proposed['descripcion_literal']):
            raise ValueError('Evidencia de significado cambió: ' + proposed['definicion_iri'])
    physical = columnas(inventory, inspection, base)
    col_by_iri = {c['iri']: c for c in physical}
    revisions = {m['columna']: m for m in semantic['correspondencias'] if m['estado'] == 'documental_revisada'}
    for proposed in propuesta('agente-4-correspondencias')['propuestas_revisadas']:
        current = variables[proposed['definicion_iri']]
        revision = revisions.get(proposed['columna'])
        if (not revision or revision['variable'] != proposed['variable_iri']
                or col_by_iri[proposed['columna']]['huella_esquema'] != proposed['huellas']['huella_esquema']
                or current['huella_fila_diccionario'] != proposed['huellas']['huella_fila_diccionario']):
            raise ValueError('Correspondencia integrada sin evidencia vigente')
    resolutions = []
    # Eight reference bindings. Sampled data are deliberately excluded: identification
    # rests on the original dictionary and the complete catalogue, not sample membership.
    for p in propuesta('agente-1-catalogos')['propuestas']:
        origin, ref = p['referencia_original'], URIRef(p['referencia_original']['iri'])
        for iri, h in zip(origin['definicion_iris'], origin['huellas_fila_diccionario'], strict=True):
            v = variables[iri]
            if v['huella_fila_diccionario'] != h or v['conjunto'] != p['conjunto']:
                raise ValueError('Definición de catálogo cambió: ' + iri)
            if (ref, PROV.wasDerivedFrom, URIRef(iri)) not in semantic_data:
                raise ValueError('Referencia no deriva de la definición declarada')
        if p['clasificacion'] == 'sin_resolver':
            g.add((ref, mex.limiteSemantico, Literal(p['justificacion'], lang='es')))
            continue
        r, ev = p['recurso'], p['evidencia']
        if r['conjunto'] != p['conjunto'] or resources[r['recurso']]['conjunto_id'] != p['conjunto']:
            raise ValueError('Catálogo fuera del conjunto')
        path = archivo(ev['archivo_local'])
        rows = leer_tabla(path, r['hoja'])
        header = rows[0]
        if header != ev['encabezado_catalogo'] or r['campo_catalogo'] not in header:
            raise ValueError('Encabezado de catálogo distinto del original')
        ci = header.index(r['campo_catalogo'])
        if len(rows) - 1 != ev['filas_datos_catalogo']:
            raise ValueError('Cantidad de entradas de catálogo cambió')
        codes = [texto(row[ci]) for row in rows[1:]]
        if len(codes) != len(set(codes)) or any(not code for code in codes):
            raise ValueError('Catálogo con claves vacías o duplicadas')
        scheme = node('published-code-reference', ev['sha256'] + '/' + r['hoja'])
        g.add((scheme, RDF.type, mex.CatalogoPublicado)).add((scheme, RDF.type, SKOS.ConceptScheme))
        g.add((scheme, RDFS.label, Literal('Catálogo publicado: ' + r['hoja'], lang='es')))
        g.add((scheme, PROV.wasDerivedFrom, URIRef(r['url'])))
        lit(scheme, 'archivoSha256', ev['sha256']); lit(scheme, 'hojaOriginal', r['hoja'])
        lit(scheme, 'estadoVigencia', 'no_documentada')
        lit(scheme, 'encabezadoOriginal', json.dumps(header, ensure_ascii=False))
        ni = header.index('Presa') if 'Presa' in header else 1
        for rowno, row in enumerate(rows[1:], 2):
            entry = node('published-code-entry', ev['sha256'] + '/' + r['hoja'] + '/' + texto(row[ci]))
            g.add((entry, RDF.type, mex.EntradaCatalogoPublicado)).add((entry, RDF.type, SKOS.Concept))
            g.add((entry, SKOS.inScheme, scheme)).add((entry, SKOS.notation, Literal(texto(row[ci]))))
            g.add((entry, RDFS.label, Literal(texto(row[ci]) + ' — ' + texto(row[ni]), lang='es')))
            g.add((entry, PROV.wasDerivedFrom, URIRef(r['url'])))
            lit(entry, 'nombreOriginal', texto(row[ni])); lit(entry, 'filaOriginal', rowno)
            lit(entry, 'filaOriginalJson', json.dumps(row, ensure_ascii=False))
        binding = node('published-code-binding', origin['id'])
        g.add((binding, RDF.type, mex.VinculoCatalogoPublicado))
        g.add((binding, RDFS.label, Literal(origin['variable_diccionario'] + ' — ' + r['hoja'], lang='es')))
        g.add((binding, mex.referenciaDeclarada, ref)).add((binding, mex.catalogoRecuperado, scheme))
        for iri in origin['definicion_iris']:
            g.add((binding, mex.variableDestino, URIRef(variables[iri]['iri'])))
            g.add((binding, PROV.wasDerivedFrom, URIRef(iri)))
        g.add((binding, PROV.wasDerivedFrom, URIRef(r['url'])))
        g.add((binding, mex.criterioRevision, Literal(r['correspondencia_campos_justificacion'], lang='es')))
        g.add((ref, mex.catalogoRecuperado, scheme))
        semantic_data.set((ref, mex.estadoReferenciaCatalogo, Literal('recurso_catalogo_identificado')))
        resolutions.append({'referencia': str(ref), 'catalogo': str(scheme), 'propuesta': p['id_propuesta']})

    # Observations within one 2020 study. Raw strings must already have exact widths;
    # unlike the exploratory agent script, no zfill, float conversion or name trimming.
    p2 = propuesta('agente-2-localidades')
    f = p2['fuentes_examinadas'][0]
    path = archivo('propuestas/agente-2-localidades/' + f['archivo'])
    rows = leer_tabla(path, 'IML_2020'); header = rows[0]; loc_records = []
    for rowno, row in enumerate(rows[1:], 2):
        d = {k: row[header.index(k)] for k in ['CVE_LOC', 'ENT', 'MUN', 'LOC', 'NOM_MUN', 'NOM_LOC']}
        for key, n in [('CVE_LOC', 9), ('ENT', 2), ('MUN', 3), ('LOC', 4)]:
            if not isinstance(d[key], str) or not re.fullmatch('[0-9]{' + str(n) + '}', d[key]):
                raise ValueError('Clave de localidad no literal: ' + str(d))
        if d['CVE_LOC'] != d['ENT'] + d['MUN'] + d['LOC'] or d['ENT'] != '26' or not d['NOM_LOC']:
            raise ValueError('Composición de localidad inconsistente')
        loc_records.append({'fila': rowno, **d})
    if len(loc_records) != 1850 or len({d['CVE_LOC'] for d in loc_records}) != 1850:
        raise ValueError('Cobertura de referencia derivada cambió')
    local_scheme = node('observed-locality-reference', f['sha256'] + '/IML_2020')
    g.add((local_scheme, RDF.type, mex.ReferenciaLocalidadesObservada)).add((local_scheme, RDF.type, SKOS.ConceptScheme))
    g.add((local_scheme, RDFS.label, Literal('Localidades observadas en marginación 2020: 1,850', lang='es')))
    g.add((local_scheme, PROV.wasDerivedFrom, URIRef(f['url'])))
    lit(local_scheme, 'archivoSha256', f['sha256']); lit(local_scheme, 'hojaOriginal', 'IML_2020')
    lit(local_scheme, 'estadoReferenciaObservada', 'derivada_observada_sin_completitud_oficial')
    lit(local_scheme, 'estadoVigencia', 'edicion_INEGI_no_cotejada')
    g.add((local_scheme, mex.limiteSemantico, Literal('Localidades incluidas en este estudio de 2020; no catálogo oficial completo ni equivalencia con una edición INEGI. No asigna identidades a otras columnas.', lang='es')))
    resource = node('resource', '8e684391-9137-4359-b654-f19602911421')
    g.add((resource, mex.referenciaLocalidadesObservada, local_scheme))
    for d in loc_records:
        entry = node('observed-locality', f['sha256'] + '/' + d['CVE_LOC'])
        g.add((entry, RDF.type, mex.LocalidadObservada)).add((entry, RDF.type, SKOS.Concept))
        g.add((entry, SKOS.inScheme, local_scheme)).add((entry, SKOS.notation, Literal(d['CVE_LOC'])))
        g.add((entry, RDFS.label, Literal(d['CVE_LOC'] + ' — ' + d['NOM_LOC'] + ' — ' + d['NOM_MUN'], lang='es')))
        g.add((entry, PROV.wasDerivedFrom, URIRef(f['url'])))
        for key, val in [('nombreOriginal', d['NOM_LOC']), ('nombreMunicipioOriginal', d['NOM_MUN']),
                         ('componenteEntidad', d['ENT']), ('componenteMunicipio', d['MUN']),
                         ('componenteLocalidad', d['LOC']), ('filaOriginal', d['fila'])]:
            lit(entry, key, val)

    territorial = json.loads((ROOT / 'data/territorio/catalogos.json').read_text(encoding='utf-8'))
    refs = {(r['recurso'], r['hoja']): r for r in territorial['referencias']}
    cache, checks = {}, []
    for check in propuesta('agente-5-aplicacion-territorial')['comprobaciones']:
        q, refinfo = check['consulta'], check['referencia_catalogo']
        path = archivo('propuestas/agente-5-aplicacion-territorial/' + q['archivo_local'])
        if path not in cache:
            cache[path] = leer_tabla(path)
        rows = cache[path]; header, values = rows[0], rows[1:]
        if header != q['header_fisico']:
            raise ValueError('Encabezado de aplicación territorial cambió')
        reference = refs[(refinfo['recurso'], refinfo['hoja'])]
        if reference['archivo_sha256'] != refinfo['sha256']:
            raise ValueError('Referencia de aplicación territorial cambió')
        cv = [v for v in semantic['variables'] if v['diccionario'] == check['diccionario'] and v['variable'] == check['variable']]
        if len(cv) != 1 or cv[0]['huella_fila_diccionario'] != check['huella_fila_diccionario']:
            raise ValueError('Definición de aplicación territorial cambió')
        keys, state = [], check['estado_resultado']
        if check['columna_fisica'] is None:
            if check['variable'] in header or check['estado_cobertura'] != 'encabezado_completo':
                raise ValueError('Ausencia de columna mal documentada')
        else:
            ci = header.index(check['columna_fisica'])
            vals = [texto(r[ci]).strip() for r in values]
            keys = sorted(set(vals) - {''})
            expected_keys = sorted(check['coincidencias']['claves'] + check['excepciones']['sin_correspondencia'])
            if keys != expected_keys or len(values) != check['cobertura']['filas'] or vals.count('') != check['excepciones']['nulos']:
                raise ValueError('Cobertura o claves de comprobación cambiaron')
            catalog_keys = {e['codigo'] for e in reference['entradas']}
            outside = sorted(set(keys) - catalog_keys)
            if outside != check['excepciones']['sin_correspondencia']:
                raise ValueError('Pertenencia a catálogo no reproducible')
            if 'CLAVE' in header and 'CLAVE_ENT' in header:
                mi, ei = header.index('CLAVE'), header.index('CLAVE_ENT')
                pairs = {(texto(r[mi]).strip(), texto(r[ei]).strip()) for r in values}
                if len(pairs) != 72 or any(not composicion_exacta(a, b) for a, b in pairs):
                    raise ValueError('Composición municipal no reproducible')
        subject = node('territorial-application-check', check['recurso_datos'] + '/' + check['variable'])
        g.add((subject, RDF.type, mex.ComprobacionAplicacionTerritorial))
        g.add((subject, RDFS.label, Literal(check['variable'] + ' — ' + check['periodo'] + ' — ' + state, lang='es')))
        g.add((subject, mex.recursoComprobado, node('resource', check['recurso_datos'])))
        g.add((subject, mex.variableDestino, URIRef(cv[0]['iri'])))
        scheme = node('territorial-reference', reference['archivo_sha256'] + '/' + reference['hoja'])
        g.add((subject, mex.referenciaComprobada, scheme))
        g.add((subject, PROV.wasDerivedFrom, URIRef(q['url'])))
        g.add((subject, PROV.wasDerivedFrom, URIRef(reference['fuente'])))
        g.add((subject, PROV.wasDerivedFrom, URIRef((ROOT / 'data/integracion/entregas/agente-5-aplicacion-territorial/propuestas.json').as_uri())))
        for key, val in [('archivoSha256', q['sha256_archivo']), ('sha256CatalogoComprobado', refinfo['sha256']),
                         ('campoComprobado', check['columna_fisica'] or check['variable']),
                         ('periodoComprobado', check['periodo']), ('estadoComprobacion', state),
                         ('coberturaComprobacion', check['estado_cobertura']), ('filasComprobadas', check['cobertura']['filas']),
                         ('evidenciaComprobacionJson', json.dumps(check, ensure_ascii=False)),
                         ('ubicacionComprobadaJson', json.dumps(check.get('columna_ubicacion'), ensure_ascii=False))]:
            lit(subject, key, val)
        g.add((subject, mex.limiteSemantico, Literal(check['limitaciones'], lang='es')))
        for key in keys:
            lit(subject, 'claveObservada', key)
        for key in check['excepciones']['especiales_no_municipales']:
            lit(subject, 'claveEspecialNoMunicipal', key)
        candidates = [c for c in physical if c['recurso'] == check['recurso_datos'] and c['campo'] == check['columna_fisica']]
        file_candidates = [c for c in candidates if '/file-column/' in c['iri']]
        candidates = file_candidates or candidates
        if len(candidates) == 1:
            g.add((subject, mex.columnaComprobada, URIRef(candidates[0]['iri'])))
        checks.append({'iri': str(subject), 'estado': state, 'columna_enlazada': len(candidates) == 1})
    summary = {'catalogos_declarados_resueltos_nuevos': len(resolutions),
               'catalogos_publicados_distintos': len(set(g.subjects(RDF.type, mex.CatalogoPublicado))),
               'entradas_catalogos_publicados': len(set(g.subjects(RDF.type, mex.EntradaCatalogoPublicado))),
               'localidades_observadas': len(loc_records), 'comprobaciones_territoriales': len(checks),
               'estados_comprobaciones': dict(Counter(c['estado'] for c in checks)),
               'comprobaciones_con_columna_del_grafo': sum(c['columna_enlazada'] for c in checks),
               'limites': ['Sin vigencia histórica acreditada', 'Sin asignación de localidades a otras fuentes',
                          'Sin cruces entre poblaciones', 'Muestras del agente 1 excluidas de las afirmaciones integradas']}
    expected = dict(config.get('esperado_actual', config['esperado']))
    r6cfg = ROOT / 'ontologia/revision-ronda-6.json'
    if r6cfg.exists():
        # R6 añade como máximo su vínculo auditado; el esperado histórico queda intacto.
        r6exp = (json.loads(r6cfg.read_text(encoding='utf-8')).get('estados_correspondencias_esperados') or {})
        if 'documental_revisada' in r6exp:
            expected['correspondencias_revisadas'] = r6exp['documental_revisada']

    r7cfg = ROOT / 'ontologia/revision-ronda-7.json'
    if r7cfg.exists():
        # R7: cadena auditada, luego deltas de backup congelado + accepted.
        # Sin defaults permisivos ni observados como esperados.
        if not r6cfg.exists():
            raise ValueError('R7 exige su base R6 instalada')

        from integrar_ronda_6 import cargar as cargar_ronda_6
        from integrar_ronda_7 import cargar as cargar_ronda_7
        r6c, r6f, r6base, r6accepted, r6s = cargar_ronda_6()
        r7c, r7f, r7base, r7accepted, r7sources = cargar_ronda_7()
        sources.update({k: v for k, v in r7sources.items() if k not in sources})
        if r7base['reglas_descripcion_exacta'] != r6base['reglas_descripcion_exacta'] + r6accepted['reglas'] \
                or r7base['vinculos_documentales_revisados'] != r6base['vinculos_documentales_revisados'] + r6accepted['vinculos']:
            raise ValueError('Base R7 discontinua respecto a cadena R6')
        for k, v in r6base['conceptos'].items():
            if r7base['conceptos'].get(k) != v:
                raise ValueError('Concepto base alterado en cadena R7: ' + k)
        for k, v in r6accepted['conceptos'].items():
            if r7base['conceptos'].get(k) != v:
                raise ValueError('Concepto R6 ausente en base R7: ' + k)
        n_reglas, n_vinc, n_con = (len(r7accepted[k]) for k in ('reglas', 'vinculos', 'conceptos'))
        if (n_reglas, n_vinc, n_con) != (31, 3, 31):
            raise ValueError('R7 fuera del lote auditado 31/3/31')
        # Pre-estado R6 desde cobertura congelada con SHA del respaldo-previo.
        bkroot = ROOT / 'propuestas/plan-paralelo-brechas-sin-gemma/ejecucion-20261024-072207Z/integracion/respaldo-previo'
        bkman = json.loads((bkroot / 'manifiesto-respaldo.json').read_text(encoding='utf-8'))
        pin = {c['original']: c['sha256'] for c in bkman['copias']}
        if sha(bkroot / 'manifiesto-respaldo.json') != 'ea354f8a635e298ab6c886ea12fe3ebf4b592658a8592514c182842f4420d5cd':
            raise ValueError('Manifiesto de respaldo-previo alterado')
        bkcov = bkroot / 'data/semantica/cobertura.json'
        if sha(bkcov) != pin.get('data/semantica/cobertura.json'):
            raise ValueError('Cobertura R6 de respaldo alterada')
        bkcob = json.loads(bkcov.read_text(encoding='utf-8'))
        pre_est = {v['definicion_iri']: v['estado'] for v in bkcob['variables']}
        backed = [r['definiciones'][0] for r in r7accepted['reglas']]
        if any(pre_est.get(d) == 'respaldada_por_definicion' for d in backed):
            raise ValueError('R7 re-respalda definición previa')
        de_cand = [d for d in backed if pre_est.get(d) == 'candidata_sin_definicion_suficiente']
        de_fuera = [d for d in backed if pre_est.get(d) == 'fuera_del_nucleo_actual']
        if (len(de_cand), len(de_fuera)) != (3, 28):
            raise ValueError('R7 sin partición 3 candidatas + 28 fuera')
        pre_var = (json.loads(r6cfg.read_text(encoding='utf-8')).get('estados_variables_esperados') or {})
        for clave in ('respaldada_por_definicion', 'fuera_del_nucleo_actual', 'candidata_sin_definicion_suficiente'):
            if clave not in pre_var:
                raise ValueError('Ancla histórica R6 ausente: ' + clave)
        if len(r7base['conceptos']) + n_con != 1106:
            raise ValueError('R7 sin 1106 conceptos derivados')
        expected['conceptos'] = len(r7base['conceptos']) + n_con
        expected['variables_respaldadas'] = pre_var['respaldada_por_definicion'] + n_reglas
        if expected['variables_respaldadas'] != 1423:
            raise ValueError('R7 sin 1423 variables derivadas')
        expected['variables_pendientes'] = pre_var['candidata_sin_definicion_suficiente'] - len(de_cand)
        if expected['variables_pendientes'] != 16:
            raise ValueError('R7 sin 16 candidatas de significado')
        expected['correspondencias_revisadas'] = len(r7base['vinculos_documentales_revisados']) + n_vinc
        if expected['correspondencias_revisadas'] != 4478:
            raise ValueError('R7 sin 4478 correspondencias derivadas')
        # 28 coincidencias automáticas nuevas: candidatas visibles, no aprobadas.
        pre_cand_cols = {m['columna'] for m in bkcob['correspondencias'] if m['estado'] == 'candidata_por_nombre'}
        promovidas = {v['columna'] for v in r7accepted['vinculos']}
        if not promovidas <= pre_cand_cols:
            raise ValueError('Vínculo R7 fuera de candidatas previas')
        viejas = pre_cand_cols - promovidas
        if len(viejas) != 17:
            raise ValueError('R7 sin 17 candidatas anteriores')
        en_vivo = {m['columna'] for m in semantic['correspondencias'] if m['estado'] == 'candidata_por_nombre'}
        nuevas = en_vivo - pre_cand_cols
        if len(nuevas) != 28 or not nuevas.isdisjoint(viejas):
            raise ValueError('R7 sin 28 coincidencias nuevas exactas')
        vardef = {v['iri']: v['definicion_iri'] for v in semantic['variables'] if 'iri' in v}
        for m in semantic['correspondencias']:
            if m['columna'] in nuevas and vardef.get(m['variable']) not in backed:
                raise ValueError('Coincidencia nueva sin respaldo R7: ' + m['columna'])
        if len(en_vivo) != 45:
            raise ValueError('R7 sin 45 candidatas (17 anteriores + 28 nuevas)')
        if bkcob['resumen']['columnas_con_correspondencia'] + len(nuevas) != 4523:
            raise ValueError('Pin histórico de columnas con correspondencia alterado')
        if semantic['resumen']['columnas_con_correspondencia'] != 4523:
            raise ValueError('R7 sin 4523 columnas con correspondencia')
        # Pendientes 140 = fuera 124 + candidatas 16, con conservación.
        fuera = pre_var['fuera_del_nucleo_actual'] - len(de_fuera)
        if fuera != 124:
            raise ValueError('R7 sin 124 fuera derivados')
        tot = semantic['resumen']['definiciones_con_nombre_examinadas']
        resp = semantic['resumen']['estados_variables']['respaldada_por_definicion']
        cand = semantic['resumen']['estados_variables']['candidata_sin_definicion_suficiente']
        if resp + fuera + cand != tot or fuera + cand != 140:
            raise ValueError('R7 sin conservación 1423+124+16')
        # Materialización fiel de cada regla aceptada (la cadena completa se
        # valida en integrar_ronda_7.construir, fase R7 de la misma corrida).
        variables = {v['definicion_iri']: v for v in semantic['variables']}
        for row in r7accepted['reglas']:
            observed = variables[row['definiciones'][0]]
            if observed.get('estado') != 'respaldada_por_definicion' \
                    or observed.get('regla') != row['id'] \
                    or observed.get('concepto') != row['concepto'] \
                    or observed.get('huella_fila_diccionario') != row['huella_fila_diccionario']:
                raise ValueError('Regla R7 no materializada: ' + row['id'])


    actual = {'conceptos': semantic['resumen']['conceptos_compartidos'],
              'variables_respaldadas': semantic['resumen']['estados_variables']['respaldada_por_definicion'],
              'variables_pendientes': semantic['resumen']['estados_variables']['candidata_sin_definicion_suficiente'],
              'correspondencias_revisadas': semantic['resumen']['estados_correspondencias']['documental_revisada'],
              'localidades_observadas': len(loc_records), 'comprobaciones': len(checks)}
    for key, val in actual.items():
        if val != expected[key]:
            raise ValueError('Resultado inesperado de integración: ' + key + ' = ' + str(val))
    summary['verificacion_conteos'] = actual
    summary['resoluciones_catalogos'] = resolutions
    summary['comprobaciones'] = checks
    summary['respaldo_previo'] = manifest['respaldo']
    summary['fuentes_fijadas'] = len({m['archivo'] for m in manifest['fuentes'].values()})
    (ROOT / 'data/integracion/integracion.ttl').write_text(g.serialize(format='turtle'), encoding='utf-8')
    (ROOT / 'reports/integracion-agentes.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    lines = ['# Integración central de las cinco entregas — ronda 1', '',
             'La tabla siguiente conserva los resultados históricos de la primera ronda. El estado vigente se documenta en reports/consolidacion-semantica.md y reports/cobertura-muse.md.', '',
             'Las carpetas originales de propuestas se conservan. La generación lee copias de evidencia fijadas por SHA-256 en data/integracion; no depende de cambios posteriores en las carpetas de agentes ni consulta la red.', '',
             '| Resultado | Cantidad |', '| --- | ---: |',
             '| Conceptos compartidos | 103 |', '| Variables con significado documental | 232 |',
             '| Variables candidatas pendientes | 23 |', '| Correspondencias revisadas (62 anteriores + 100) | 162 |',
             '| Correspondencias candidatas por nombre | 792 |', '| Columnas ambiguas / vínculos ambiguos | 3 / 6 |',
             '| Referencias a catálogos identificadas (8 anteriores + 8) | 16 |', '| Referencias a catálogos sin resolver | 1 |',
             '| Localidades observadas en IML 2020 | 1850 |', '| Comprobaciones territoriales | 28 |', '',
             '## Decisiones y límites', '',
             '- Agente 1: ocho vínculos documentales a hojas completas. Las cuatro publicaciones educativas conservan sus referencias por conjunto, aunque comparten bytes. Las muestras de datos no se integran: algunos metadatos de muestra del JSON repiten el campo de servicios en otras variables; no se usan como evidencia de pertenencia. La identificación se sustenta en diccionario y catálogo completo. La referencia des_mpo_res sigue sin resolver y conserva las dos definiciones de origen.',
             '- Agente 2: 1850 observaciones con composición cruda exacta 2+3+4, sin padding ni conversión numérica. Son localidades incluidas en IML_2020, no un catálogo oficial completo ni una equivalencia con una edición INEGI. No se importó el contraste externo puntual como acreditación de cobertura.',
             '- Agente 3: 20 conceptos/reglas, restringidos simultáneamente por definición, diccionario, conjunto, nombre y descripción. Esta ronda promocionó 24 significados; entonces quedaron 23 candidatos y 1308 definiciones fuera del núcleo. Las ampliaciones posteriores se registran en la cobertura vigente.',
             '- Agente 4: 99 correspondencias y un alias técnico limitado a la columna de divorcios 2023. a–o de nacimientos 2024 no se promueve. Las 62 revisiones existentes se conservan.',
             '- Agente 5: 16 comprobaciones compatibles, 11 con excepciones y una ausencia de columna. Se vuelven a leer las 25 fuentes y se verifican claves, filas, pertenencia a catálogo y composición exacta. No se asigna un sistema de identificadores a columnas ni se promociona su correspondencia semántica. 26998/26999 siguen siendo categorías no municipales. Víctimas 2015 solo tiene revisión completa de encabezado; validación de Clave_Ent no aplicable.', '',
             'No se acredita vigencia histórica, no se ejecutan nuevos cruces ni se crean owl:sameAs/skos:exactMatch. Las 5 inconsistencias documentales previas siguen visibles, incluida la limitación del ejemplo de formato municipal.', '',
             '## Evidencia y reproducción', '',
             '- ontologia/revision-agentes.json fija el manifiesto y los conteos de la integración aprobada.',
             '- data/integracion/manifiesto.json conserva las huellas de las cinco entregas y 37 archivos fuente únicos.',
             '- scripts/ampliar_portal.py --load regenera las capas y valida SHACL antes de sincronizar Studio.',
             '- scripts/sincronizar_studio_portal.py respalda el grafo vivo y conserva las ediciones ajenas a la ampliación.', '',
             'Respaldo previo de archivos centrales: ' + manifest['respaldo'], '']
    (ROOT / 'reports/integracion-agentes.md').write_text('\n'.join(lines), encoding='utf-8')
    return g, model, summary, sources
