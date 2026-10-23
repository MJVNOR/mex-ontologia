"""Contraste IML_2020 (evidencia central) vs ITER Censo 2020 Sonora.

Uso:
  uv run python propuestas/ronda-2/agente-3-localidades-inegi/contrastar.py          # recalcula y escribe entrega
  uv run python propuestas/ronda-2/agente-3-localidades-inegi/contrastar.py --check  # solo verifica, no escribe; sale !=0 si algo difiere

Reglas: codigos originales intactos; padding solo si el original textual es ^[0-9]+$
(ASCII) y con ancho documentado; cada transformacion queda registrada aparte.
En estas dos fuentes no hace falta padding (0 transformaciones).
"""
import csv
import hashlib
import json
import os
import re
import sys
import unicodedata
import zipfile
from collections import Counter, defaultdict

from openpyxl import load_workbook

OUT = os.path.join("propuestas", "ronda-2", "agente-3-localidades-inegi")
F = os.path.join(OUT, "fuentes")
ZP = os.path.join(F, "iter_26_cpv2020_csv.zip")
DATASET = "iter_26_cpv2020/conjunto_de_datos/conjunto_de_datos_iter_26CSV20.csv"
IML_PATH = os.path.join("data", "integracion", "fuentes",
                         "771cc22d44fcba3d70c8762d6c90aebfbe1dd738e5f94cf3119d80d3fc01773e.xlsx")
ENTREGA = os.path.join(OUT, "contraste-localidades.json")
SOLO_ITER = os.path.join(OUT, "solo-ITER-CVE.json")
FUENTE_JSON = os.path.join(OUT, "inegi-fuente.json")

ASCII_DIGITOS = re.compile(r"[0-9]+\Z")  # [0-9] es ASCII por definicion; \d aceptaria digitos Unicode


def sha256(p):
    h = hashlib.sha256()
    with open(p, "rb") as fh:
        for b in iter(lambda: fh.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def norm_nombre(s):
    s = (s or "").strip().lower()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s)


def tramo_original(raw):
    """Texto original estricto: solo str (recortado en bordes) o int exacto.
    float/decimal/None/otros -> None (no se trunca ni se sanea)."""
    if raw is None:
        return None
    if isinstance(raw, bool):
        return None
    if isinstance(raw, int):
        return str(raw)
    if isinstance(raw, str):
        return raw.strip()
    return None


def validar_tramo(texto, ancho, campo, ref, errores):
    """El original debe ser digitos ASCII 0-9 y medir <= ancho. Devuelve (valor, necesito_pad)."""
    if texto is None or ASCII_DIGITOS.fullmatch(texto) is None:
        errores.append({"campo": campo, "original": texto, "motivo": "no-es-digito-ASCII", "ref": ref})
        return None, False
    if len(texto) > ancho:
        errores.append({"campo": campo, "original": texto, "motivo": f"excede-ancho-{ancho}", "ref": ref})
        return None, False
    if len(texto) < ancho:
        return texto.zfill(ancho), True
    return texto, False


def leer_iter(errores):
    with zipfile.ZipFile(ZP) as z:
        crudo = z.read(DATASET)
    if crudo.startswith(b"\xef\xbb\xbf"):
        crudo = crudo[3:]
    try:
        texto = crudo.decode("utf-8")
    except UnicodeDecodeError:
        texto = crudo.decode("cp1252")
    lineas = texto.splitlines()
    encabezado_crudo = lineas[0]
    # Parseo con quotechar estandar (comillas dobles). Las comillas simples NO son
    # delimitador: usar quotechar="'" rompe 2 filas (ver nota en formato_csv).
    rd = csv.DictReader(lineas)
    campos = rd.fieldnames or []
    filas, malformadas = [], []
    for n, r in enumerate(rd, start=2):  # n = numero de linea fisica (1 = encabezado)
        if r is None or None in r or len(r) != len(campos):
            malformadas.append(n)
            continue
        filas.append(r)
    fmt = {
        "codificacion": "utf-8 con BOM (respaldo cp1252)",
        "separador": "coma",
        "quotechar_usado": 'comillas dobles (estandar)',
        "lineas_totales": len(lineas),
        "columnas": len(campos),
        "encabezado_crudo": encabezado_crudo[:200],
        "filas_leidas": len(filas),
        "filas_malformadas": malformadas,
        "nota_comillas_simples": "El archivo no usa comillas simples como delimitador ni como envoltura "
                                 "de valores; parsear con quotechar=\"'\" produce 2 filas malformadas "
                                 "(lineas 1197 y 7495: NOM_LOC con comillas dobles literales, "
                                 "Caborca LOC 1111 y San Ignacio Rio Muerto LOC 0337). No se retira "
                                 "ninguna envoltura: los valores comparados son los literales del CSV.",
    }
    return filas, campos, fmt


def leer_iml(errores, transformaciones):
    wb = load_workbook(IML_PATH, read_only=True, data_only=True)
    ws = wb["IML_2020"]
    filas = list(ws.iter_rows(values_only=True))
    hdr = [str(v).strip() if v is not None else "" for v in filas[0]]
    idx = {h: i for i, h in enumerate(hdr)}
    recs = []
    for n, r in enumerate(filas[1:], start=2):
        g = lambda c: r[idx[c]]  # noqa: E731
        ref = f"IML_2020 fila {n}"
        t = {c: tramo_original(g(c)) for c in ("CVE_LOC", "ENT", "MUN", "LOC")}
        v = {}
        for c, ancho in (("ENT", 2), ("MUN", 3), ("LOC", 4), ("CVE_LOC", 9)):
            val, pad = validar_tramo(t[c], ancho, c, ref, errores)
            v[c] = val
            if pad:
                transformaciones.append({"campo": c, "original": t[c], "resultado": val,
                                         "regla": f"zfill({ancho})", "ref": ref})
        cve = v["CVE_LOC"]
        if cve is not None and cve != (v["ENT"] or "") + (v["MUN"] or "") + (v["LOC"] or ""):
            errores.append({"campo": "CVE_LOC", "original": t["CVE_LOC"],
                            "motivo": "composicion-9!=2+3+4", "ref": ref})
            cve = None
        recs.append({"CVE": cve, "POB_TOT": g("POB_TOT"),
                     "NOM_MUN": str(g("NOM_MUN") or "").strip(),
                     "NOM_LOC": str(g("NOM_LOC") or "").strip()})
    wb.close()
    return recs, hdr


def calcular():
    errores, transformaciones = [], []
    filas_iter, campos, fmt = leer_iter(errores)
    recs_iml, hdr_iml = leer_iml(errores, transformaciones)

    agregados, especiales, localidades = [], [], []
    for r in filas_iter:
        ent, mun, loc = r["ENTIDAD"], r["MUN"], r["LOC"]
        for c, v, ancho in (("ENTIDAD", ent, 2), ("MUN", mun, 3), ("LOC", loc, 4)):
            t = tramo_original(v)
            ok, _ = validar_tramo(t, ancho, f"ITER.{c}", f"ITER {ent}|{mun}|{loc}", errores)
            if ok is None:
                ent = mun = loc = None
                break
        else:
            cve = ent + mun + loc
        if ent is None:
            continue
        if mun == "000" or loc == "0000":
            agregados.append((cve, r))
        elif loc in ("9998", "9999"):
            especiales.append((cve, r))
        else:
            localidades.append((cve, r))

    cves_iter = Counter(c for c, _ in localidades)
    duplicadas_iter = sorted(c for c, n in cves_iter.items() if n > 1)
    if duplicadas_iter:
        errores.append({"campo": "ITER.CVE", "motivo": "duplicadas", "cves": duplicadas_iter[:10]})
    dic_iter = {c: {"NOM_MUN": r["NOM_MUN"].strip(), "NOM_LOC": r["NOM_LOC"].strip(),
                    "POBTOT": r["POBTOT"].strip()} for c, r in localidades}

    vistos, duplicadas_iml, dic_iml, sin_clave = set(), [], {}, 0
    for r in recs_iml:
        if r["CVE"] is None:
            sin_clave += 1
            continue
        if r["CVE"] in vistos:
            duplicadas_iml.append(r["CVE"])
        vistos.add(r["CVE"])
        dic_iml[r["CVE"]] = {"NOM_MUN": r["NOM_MUN"], "NOM_LOC": r["NOM_LOC"]}
    if duplicadas_iml:
        errores.append({"campo": "IML.CVE_LOC", "motivo": "duplicadas", "cves": duplicadas_iml[:10]})
    esp_iml = [c for c in dic_iml if c[2:5] == "000" or c[5:] in ("0000", "9998", "9999") or c.endswith("9999")]

    s_iml, s_iter = set(dic_iml), set(dic_iter)
    inter = s_iml & s_iter
    solo_iml = sorted(s_iml - s_iter)
    solo_iter = sorted(s_iter - s_iml)
    dif_exact, dif_norm = [], []
    for c in sorted(inter):
        a, b = dic_iml[c], dic_iter[c]
        if (a["NOM_MUN"], a["NOM_LOC"]) != (b["NOM_MUN"], b["NOM_LOC"]):
            dif_exact.append({"CVE": c, "IML": a, "ITER": b})
        if norm_nombre(a["NOM_MUN"]) != norm_nombre(b["NOM_MUN"]) or \
           norm_nombre(a["NOM_LOC"]) != norm_nombre(b["NOM_LOC"]):
            dif_norm.append({"CVE": c, "IML": a, "ITER": b})

    def toint(x):
        try:
            return int(str(x).strip())
        except (ValueError, AttributeError):
            return None

    pob_igual, pob_dist, iml_nan, iter_nan = 0, [], 0, 0
    for c in sorted(inter):
        a = toint(next(r["POB_TOT"] for r in recs_iml if r["CVE"] == c))
        b = toint(dic_iter[c]["POBTOT"])
        if a is None:
            iml_nan += 1
        elif b is None:
            iter_nan += 1
        elif a == b:
            pob_igual += 1
        else:
            pob_dist.append(c)

    cub = Counter(c[:5] for c in s_iter)
    for c in s_iml:
        cub.setdefault(c[:5], cub.get(c[:5], 0))
    cov = [{"CVE_ENT_MUN": k, "IML": sum(1 for c in s_iml if c[:5] == k),
            "ITER_loc": sum(1 for c in s_iter if c[:5] == k)} for k in sorted(
                {c[:5] for c in s_iml} | {c[:5] for c in s_iter})]

    dist = lambda universo: dict(sorted(Counter(
        next((f"[{a},{b})" for a, b in
              [(0, 1), (1, 2), (2, 5), (5, 10), (10, 50), (50, 100), (100, 500),
               (500, 1000), (1000, 2500), (2500, 10 ** 9)] if a <= v < b), "?")
        for v in (toint(dic_iter[c]["POBTOT"]) for c in universo) if v is not None).items()))

    res = {
        "fuente_iml": {"ruta_evidencia_central": IML_PATH, "sha256": sha256(IML_PATH),
                        "hoja": "IML_2020", "filas_datos": len(recs_iml),
                        "encabezado": hdr_iml},
        "fuente_iter": {"archivo_zip": os.path.basename(ZP), "sha256": sha256(ZP),
                        "dataset": DATASET, "filas_totales_csv": len(filas_iter),
                        "agregados_MUN000_o_LOC0000": len(agregados),
                        "especiales_confidenciales_LOC9998_9999": len(especiales),
                        "localidades_censales_publicadas": len(localidades),
                        "claves_duplicadas_iter": duplicadas_iter},
        "formato_csv_iter": fmt,
        "iml": {"filas": len(recs_iml), "cve_unicos": len(dic_iml),
                "sin_clave_utilizable": sin_clave, "claves_duplicadas": duplicadas_iml,
                "especiales_LOC0000_9998_9999": len(esp_iml),
                "municipios": len({c[:5] for c in s_iml})},
        "contraste_localidades": {"interseccion": len(inter), "solo_IML": len(solo_iml),
                                  "solo_ITER": len(solo_iter),
                                  "mismo_CVE_distinto_nombre_literal": len(dif_exact),
                                  "mismo_CVE_distinto_nombre_normalizado": len(dif_norm)},
        "solo_IML_CVE": solo_iml,
        "mismo_CVE_distinto_nombre": dif_exact[:50],
        "poblacion_interseccion": {"igual_POBTOT": pob_igual, "distinta_POBTOT": len(pob_dist),
                                   "cves_distintos": pob_dist[:30],
                                   "iml_no_numerico": iml_nan, "iter_no_numerico": iter_nan},
        "distribucion_POBTOT_ITER_interseccion": dist(inter),
        "distribucion_POBTOT_ITER_solo_ITER": dist(solo_iter),
        "cobertura_por_municipio_ENT_MUN": cov,
        "transformaciones_registradas": transformaciones,
        "nota_transformaciones": "zfill(2/3/4/9) solo si el original textual cumple ^[0-9]+$ ASCII "
                                 "([0-9] no incluye digitos Unicode) y mide menos que el ancho; "
                                 "floats y decimales se rechazan sin truncar; originales intactos.",
        "errores_validacion": errores,
    }
    detalle_solo = [{"CVE": c, **dic_iter[c]} for c in solo_iter]
    return res, detalle_solo


def main():
    solo_check = "--check" in sys.argv
    res, detalle_solo = calcular()
    if solo_check:
        fallas = list(res["errores_validacion"])
        try:
            ent = json.load(open(ENTREGA, encoding="utf-8"))
            det = json.load(open(SOLO_ITER, encoding="utf-8"))
            fue = json.load(open(FUENTE_JSON, encoding="utf-8"))
        except (OSError, ValueError) as e:
            print(f"FALLA verificador: no se pudo leer entrega: {e}")
            return 1
        def cmp(path, a, b):
            if a != b:
                fallas.append({"campo": path, "motivo": "no-coincide-con-entrega",
                               "calculado": a, "entregado": b})
        cmp("fuente_iter.sha256", res["fuente_iter"]["sha256"], ent["fuente_iter"]["sha256"])
        cmp("fuente_iml.sha256", res["fuente_iml"]["sha256"], ent["fuente_iml"]["sha256"])
        if res["fuente_iter"]["sha256"] != fue.get("sha256_zip"):
            fallas.append({"campo": "inegi-fuente.json:sha256_zip", "motivo": "no-coincide-con-entrega",
                           "calculado": res["fuente_iter"]["sha256"],
                           "entregado": fue.get("sha256_zip")})
        for k in ("interseccion", "solo_IML", "solo_ITER", "mismo_CVE_distinto_nombre_literal",
                  "mismo_CVE_distinto_nombre_normalizado"):
            cmp(f"contraste_localidades.{k}", res["contraste_localidades"][k],
                ent["contraste_localidades"][k])
        cmp("solo_IML_CVE", res["solo_IML_CVE"], ent["solo_IML_CVE"])
        cmp("solo-ITER-CVE.json", detalle_solo, det)
        cmp("poblacion_interseccion", res["poblacion_interseccion"], ent["poblacion_interseccion"])
        cmp("filas_totales_csv", res["fuente_iter"]["filas_totales_csv"],
            ent["fuente_iter"]["filas_totales_csv"])
        if res["formato_csv_iter"]["filas_malformadas"]:
            fallas.append({"campo": "formato_csv_iter", "motivo": "filas-malformadas",
                           "lineas": res["formato_csv_iter"]["filas_malformadas"]})
        if fallas:
            print(f"VERIFICACION FALLIDA ({len(fallas)}):")
            for f_ in fallas:
                print(" -", json.dumps(f_, ensure_ascii=False)[:400])
            return 1
        print(f"VERIFICACION OK: interseccion={res['contraste_localidades']['interseccion']} "
              f"solo_IML={res['contraste_localidades']['solo_IML']} "
              f"solo_ITER={res['contraste_localidades']['solo_ITER']} "
              f"dif_nombre={res['contraste_localidades']['mismo_CVE_distinto_nombre_literal']} "
              f"pob_igual={res['poblacion_interseccion']['igual_POBTOT']} "
              f"transformaciones={len(res['transformaciones_registradas'])} "
              f"errores_validacion={len(res['errores_validacion'])}")
        return 0
    if res["errores_validacion"]:
        print("ERRORES, no se escribe entrega:")
        for f_ in res["errores_validacion"]:
            print(" -", json.dumps(f_, ensure_ascii=False)[:300])
        return 1
    with open(ENTREGA, "w", encoding="utf-8") as fh:
        json.dump(res, fh, indent=2, ensure_ascii=False)
    with open(SOLO_ITER, "w", encoding="utf-8") as fh:
        json.dump(detalle_solo, fh, ensure_ascii=False)
    print(f"OK interseccion={res['contraste_localidades']['interseccion']} "
          f"solo_IML={res['contraste_localidades']['solo_IML']} "
          f"solo_ITER={res['contraste_localidades']['solo_ITER']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
