"""Comprueba propuestas.json contra los archivos descargados. Falla si algo difiere."""
import json, os, re, hashlib, shutil
import openpyxl

B = 'propuestas/agente-1-catalogos'
doc = json.load(open(f'{B}/propuestas.json', encoding='utf-8'))
man = json.load(open(f'{B}/manifiesto_descargas.json', encoding='utf-8'))
cat = json.load(open(f'{B}/catalogos_contenido.json', encoding='utf-8'))
P = doc['propuestas']
assert len(P) == 9, f'esperaba 9, hay {len(P)}'
assert [p['id_propuesta'] for p in P] == [f'P{i:02d}' for i in range(1, 10)]
cl = [p['clasificacion'] for p in P]
assert cl.count('recurso_identificado') == 8, cl
assert cl == ['recurso_identificado'] * 8 + ['sin_resolver'], cl

def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for b in iter(lambda: f.read(1 << 20), b''):
            h.update(b)
    return h.hexdigest()

def filas_bin(path, hoja):
    tmp = path + '.check.xlsx'
    shutil.copyfile(path, tmp)
    try:
        wb = openpyxl.load_workbook(tmp, data_only=True, read_only=True)
        assert hoja in wb.sheetnames, (path, wb.sheetnames)
        ws = wb[hoja]
        rows = [[('' if v is None else str(v)) for v in r]
                for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    finally:
        os.remove(tmp)

# 1) SHA-256 de los 12 archivos: recalcular y comparar con manifiesto y propuestas
visto = set()
for a in man['archivos']:
    assert re.fullmatch(r'[0-9a-f]{64}', a.get('sha256') or ''), a['recurso']
    assert os.path.exists(a['archivo_local']), a['archivo_local']
    assert sha256(a['archivo_local']) == a['sha256'], a['archivo_local']
    visto.add((a['conjunto'], a['recurso']))
assert len(man['archivos']) == 12, len(man['archivos'])

for p in P:
    ro = p['referencia_original']
    assert ro['iri'].startswith('https://example.org/mex-ontologi/declared-code-reference/')
    assert p.get('no_aplicado') is True
    for k in ('justificacion', 'alcance', 'limitaciones', 'cambios_sugeridos_integrador',
              'evidencia', 'recurso'):
        assert p[k], f"{p['id_propuesta']} sin {k}"
    assert 'normalizacion' in p['recurso'], f"{p['id_propuesta']} sin normalizacion"
    # Sin restos sin acento donde el original los lleva
    blob = json.dumps(p, ensure_ascii=False)
    for malo in ('Formacion Academica', 'Catalogo_estatal', 'Clasificacion mexicana'):
        assert malo not in blob, f"{p['id_propuesta']} contiene {malo!r} sin acentos"
    ev = p['evidencia']
    if p['clasificacion'] == 'recurso_identificado':
        assert (ev['fuente_recurso_url'] or '').startswith('https://datos.sonora.gob.mx/dataset/')
        assert ev['archivo_local'].startswith(B + '/fuentes/')
        assert sha256(ev['archivo_local']) == ev['sha256'], p['id_propuesta']
        assert (ev['archivo_local'],) and os.path.exists(ev['archivo_local'])
        assert p['recurso']['hoja'], p['id_propuesta']

# 2) Hojas, encabezados y entradas completas contra los .bin descargados
for p in P[:8]:
    cj, rc, hoja = p['conjunto'], p['recurso']['recurso'], p['recurso']['hoja']
    rows = filas_bin(p['evidencia']['archivo_local'], hoja)
    while rows and all(c == '' for c in rows[-1]):
        rows.pop()
    # encabezado exacto del catálogo contra el archivo
    enc_prop = None
    for k, v in cat.items():
        if cj in k and rc in k:
            for h in v['hojas']:
                if h['hoja'] == hoja:
                    enc_prop = h['encabezado']
    assert enc_prop == rows[0][:len(enc_prop)], (p['id_propuesta'], rows[0])
    # entradas completas
    if 'entradas_completas' in p['evidencia']:
        assert [r[:2] for r in rows[1:]] == p['evidencia']['entradas_completas'], p['id_propuesta']
    if 'codigos' in p['evidencia']:
        assert [r[0] for r in rows[1:]] == p['evidencia']['codigos'], p['id_propuesta']
    cods = [r[0] for r in rows[1:]]
    assert len(set(cods)) == len(cods), p['id_propuesta']

# 3) P09: dos evidencias distintas, sin hoja
evs = P[8]['referencia_original']
assert evs['definicion_iris'][0].endswith('%2F177') and evs['definicion_iris'][1].endswith('%2F178')
assert evs['huellas_fila_diccionario'][0] != evs['huellas_fila_diccionario'][1]
assert P[8]['recurso']['hoja'] is None

# 4) 4 educativos: mismo SHA, distinto conjunto
assert len({p['evidencia']['sha256'] for p in P[0:4]}) == 1
assert len({p['conjunto'] for p in P[0:4]}) == 4

md = open(f'{B}/resultado.md', encoding='utf-8').read()
for s in ('P09', 'des_mpo_res', '177', '178', '53bd5e97', 'Catálogo_estatal',
          'Campo de Formación Académica', 'Educación', 'no aplicadas'):
    assert s in md, s
assert 'Formacion Academica' not in md and 'Catalogo_estatal' not in md
print('comprueba OK: literales, SHA-256, hojas, encabezados y entradas verificados contra archivos')
